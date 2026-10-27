"""第 6 章 · 三类质检 + 两类加分检查

每个 check 签名统一为 fn(ctx, params) -> {"ok": bool, "msg": str, ...detail}
这样 qc_rules.yaml 里写什么类型就跑什么检查，加规则不用改代码。
"""
import re

from docx import Document

# ---------- 检查实现 ----------


def check_required_sections(ctx: dict, params: dict) -> dict:
    """结构完整性：四个必备章节标题都在正文里"""
    sections = params.get("sections", [])
    text = ctx["text"]
    missing = [s for s in sections if s not in text]
    return {
        "ok": not missing,
        "missing": missing,
        "msg": "结构完整" if not missing else f"缺少章节：{missing}",
    }


def check_number_consistency(ctx: dict, params: dict) -> dict:
    """数据一致性：正文里每个"环比 x%"都必须在数字白名单里

    这是最像"AI 编数字"的漏洞，也是本题最值钱的一条检查。
    """
    pattern = params.get("pattern", r"环比[+-]?\d+(?:\.\d+)?%")
    analyze = ctx.get("analyze_result") or {}
    allowed = analyze.get("mom_percentages")
    if allowed is None:
        allowed = analyze.get("allowed_percentages", [])
    allowed = [float(v) for v in allowed]

    found = re.findall(pattern, ctx["text"])
    bad, checked = [], []
    for item in found:
        match = re.search(r"[-+]?\d+(?:\.\d+)?", item)
        if not match:
            continue
        num = float(match.group())
        checked.append(num)
        if not any(abs(num - a) <= 0.05 for a in allowed):
            bad.append(item)
    return {
        "ok": not bad,
        "checked": len(found),
        "allowed": allowed,
        "bad": bad,
        "msg": (f"检查 {len(found)} 处百分比，全部有出处" if not bad
                else f"这些数字查无出处：{bad}"),
    }


def check_word_count(ctx: dict, params: dict) -> dict:
    """篇幅：去掉空白后的字数落在区间内"""
    lo, hi = int(params.get("min", 0)), int(params.get("max", 10 ** 9))
    count = len(re.sub(r"\s", "", ctx["text"]))
    return {
        "ok": lo <= count <= hi,
        "count": count,
        "msg": f"字数 {count}" + ("，合格" if lo <= count <= hi else f"，不在 {lo}-{hi} 区间"),
    }


def check_sensitive(ctx: dict, params: dict) -> dict:
    """敏感信息：手机号 / 身份证 / 薪资等，命中直接拦截不自动修"""
    hits: list[str] = []
    for pattern in params.get("patterns", []):
        hits += re.findall(pattern, ctx["text"])
    return {
        "ok": not hits,
        "hits": hits,
        "msg": "无敏感信息" if not hits else f"命中敏感信息：{hits}",
    }


def check_has_image(ctx: dict, params: dict) -> dict:
    """图表存在性：docx 内嵌图片数量"""
    need = int(params.get("min", 1))
    try:
        count = len(Document(ctx["docx_path"]).inline_shapes)
    except Exception as exc:
        return {"ok": False, "count": 0, "msg": f"无法读取文档图片：{exc}"}
    return {
        "ok": count >= need,
        "count": count,
        "msg": f"文档含 {count} 张图" + ("，达标" if count >= need else f"，至少要 {need} 张"),
    }


CHECKS = {
    "required_sections": check_required_sections,
    "number_consistency": check_number_consistency,
    "word_count": check_word_count,
    "regex_blacklist": check_sensitive,
    "has_image": check_has_image,
}


def load_docx_text(path: str) -> str:
    """把 docx 正文 + 表格拍平成纯文本供检查"""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def run_check(rule: dict, ctx: dict) -> dict:
    """按 yaml 里的 type 分发，未知类型降级成 warning（不拦截、不崩溃）"""
    fn = CHECKS.get(rule["type"])
    base = {"id": rule["id"], "name": rule.get("name", rule["id"]), "type": rule["type"]}
    if fn is None:
        return {**base, "ok": False, "severity": "warning",
                "msg": f"未知的检查类型 {rule['type']}，已跳过（请检查 qc_rules.yaml）"}
    result = fn(ctx, rule.get("params", {}))
    result.update(base)
    result["severity"] = rule.get("severity", "blocking")
    return result
