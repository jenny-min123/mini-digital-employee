"""测试公共设施：造隔离环境 + 造 fixture 文件

关键点：先把 MDW_DB_PATH 指到临时目录，再 import core.db，
否则测试会往项目根的 tasks.db 里写脏数据。
"""
import os
import sys
import tempfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

_TMP_ROOT = tempfile.mkdtemp(prefix="mdw-test-")
os.environ["MDW_DB_PATH"] = os.path.join(_TMP_ROOT, "tasks.db")
os.environ.pop("MDW_STRICT_LLM", None)

import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402

CATEGORIES = ["生鲜", "日用百货", "饮料", "休闲零食"]


def _write_sales(path: str, rows: list[dict]) -> str:
    pd.DataFrame(rows).to_excel(path, index=False)
    return path


def _week_rows(week_start: str, stores: dict, days: int = 7, category_share: float = 0.25) -> list[dict]:
    rows = []
    start = pd.Timestamp(week_start)
    for store, amount in stores.items():
        for d in range(days):
            for category in CATEGORIES:
                rows.append({
                    "日期": (start + pd.Timedelta(days=d)).strftime("%Y-%m-%d"),
                    "门店": store,
                    "品类": category,
                    "销售额": round(amount * category_share, 2),
                    "订单数": max(1, int(amount * category_share / 60)),
                    "区域": "华东区" if "上海" in store else "华南区",
                })
    return rows


@pytest.fixture(scope="session")
def fixtures():
    base = os.path.join(_TMP_ROOT, "fixtures")
    os.makedirs(base, exist_ok=True)
    paths: dict[str, str] = {}

    # 正常数据：W43 正常，W44 上海浦东店下滑约 -25%
    rows = _week_rows("2026-10-19", {"深圳南山店": 40000, "上海静安店": 36000, "上海浦东店": 32000})
    rows += _week_rows("2026-10-26", {"深圳南山店": 41000, "上海静安店": 35000, "上海浦东店": 24000})
    paths["good"] = _write_sales(os.path.join(base, "good.xlsx"), rows)

    # 只有 1 天数据（T05：数据周期不足，不给环比结论）
    paths["single_day"] = _write_sales(
        os.path.join(base, "single_day.xlsx"),
        _week_rows("2026-10-26", {"深圳南山店": 40000}, days=1))

    # 销售额全为 0（T04：不得编造增长率）
    paths["zero_sales"] = _write_sales(
        os.path.join(base, "zero_sales.xlsx"),
        _week_rows("2026-10-19", {"深圳南山店": 0}) + _week_rows("2026-10-26", {"深圳南山店": 0}))

    # 缺"销售额"列（T08）
    rows_bad = _week_rows("2026-10-26", {"深圳南山店": 40000})
    for row in rows_bad:
        row.pop("销售额")
    paths["missing_col"] = _write_sales(os.path.join(base, "missing_col.xlsx"), rows_bad)

    # 列名用业务口语，验证 COLUMN_MAP
    rows_alias = _week_rows("2026-10-19", {"深圳南山店": 40000})
    alias = []
    for row in rows_alias:
        alias.append({"销售日期": row["日期"], "门店名称": row["门店"], "类别": row["品类"],
                      "销售金额": row["销售额"], "订单量": row["订单数"]})
    paths["alias_col"] = _write_sales(os.path.join(base, "alias_col.xlsx"), alias)

    # 空表（只有表头，没有数据行）
    empty_path = os.path.join(base, "empty.xlsx")
    pd.DataFrame(columns=["日期", "门店", "品类", "销售额", "订单数"]).to_excel(empty_path, index=False)
    paths["empty"] = empty_path

    # 乱码 CSV（T09）
    garbled = os.path.join(base, "garbled.csv")
    with open(garbled, "wb") as f:
        f.write(b"\x80\x81\x82\x83\xff\xfe\x00\x01\x02")
    paths["garbled_csv"] = garbled

    # 周会纪要
    minutes = os.path.join(base, "minutes.docx")
    doc = Document()
    doc.add_heading("周会纪要", level=1)
    doc.add_paragraph("华东区负责人说明：上海浦东店因商场改造客流下滑，晚间排班不足，销售额明显下降。")
    doc.add_paragraph("华南区：深圳南山店表现稳定。")
    doc.save(minutes)
    paths["minutes"] = minutes

    # 损坏的 docx
    broken = os.path.join(base, "broken.docx")
    with open(broken, "wb") as f:
        f.write(b"this is not a real docx")
    paths["broken_docx"] = broken

    return paths


@pytest.fixture()
def workdir(tmp_path, monkeypatch):
    """每个用例独立的 cwd + 独立的 output 目录"""
    monkeypatch.chdir(tmp_path)
    from core import orchestrator

    monkeypatch.setattr(orchestrator, "OUTPUT_ROOT", str(tmp_path / "output"))
    return tmp_path
