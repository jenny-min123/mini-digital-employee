"""第 5 章 · 工具三：读周会纪要、生成 Word

write_docx 必须幂等：固定输出路径覆盖写，重试不会产生"周报(1).docx"。
这一条正好命中加分项"重试时不产生重复结果"。
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from tools.base import ToolError


def read_minutes(file_path: str) -> dict:
    """读周会纪要（docx），返回纯文本"""
    if not os.path.exists(file_path):
        raise ToolError(f"文件不存在：{file_path}")
    if os.path.splitext(file_path)[1].lower() != ".docx":
        raise ToolError(f"周会纪要只支持 .docx，收到：{os.path.splitext(file_path)[1]}")

    try:
        doc = Document(file_path)
    except Exception as exc:
        raise ToolError(f"纪要文件损坏或不是合法 docx：{exc}") from exc

    chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    text = "\n".join(chunks)
    if not text:
        raise ToolError("纪要文件是空的，没有任何文字内容")
    return {"text": text, "chars": len(text), "paragraphs": len(chunks)}


def write_docx(content: dict, out_path: str) -> dict:
    """把结构化内容写成 Word。

    content = {
        "title": str, "period": str, "generated_at": str,
        "sections": [{"heading": str, "paragraphs": [str], "image": path|None}]
    }
    out_path 由调用方固定（output/{task_id}/周报.docx），重试覆盖同一路径。
    """
    if not content.get("sections"):
        raise ToolError("内容为空：sections 至少要有一节")
    if not out_path.lower().endswith(".docx"):
        raise ToolError(f"输出路径必须以 .docx 结尾：{out_path}")

    try:
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    except OSError as exc:
        raise ToolError(f"无法创建输出目录：{exc}") from exc

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    title = doc.add_heading(content.get("title", "经营周报"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"统计周期：{content.get('period', '-')}　|　生成时间：{content.get('generated_at', '-')}"
    )
    run.font.size = Pt(9)

    images = 0
    for section in content["sections"]:
        doc.add_heading(section.get("heading", ""), level=1)
        for para in section.get("paragraphs", []):
            if para and str(para).strip():
                doc.add_paragraph(str(para).strip())
        image = section.get("image")
        if image:
            if not os.path.exists(image):
                raise ToolError(f"图片不存在：{image}")
            doc.add_picture(image, width=Cm(15))
            images += 1

    try:
        doc.save(out_path)
    except PermissionError as exc:
        # 典型场景：文件正在 Word 里打开着
        raise ToolError(f"无法写入 {out_path}，请先关闭已打开的同名文档：{exc}") from exc
    except OSError as exc:
        raise ToolError(f"磁盘写入失败：{exc}") from exc

    if not os.path.exists(out_path) or os.path.getsize(out_path) == 0:
        raise ToolError(f"文件写入后为空：{out_path}")

    return {"file_path": out_path, "sections": len(content["sections"]), "images": images}


def docx_text(path: str) -> str:
    """质检用：把 docx 正文和表格都拍平成纯文本"""
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)
